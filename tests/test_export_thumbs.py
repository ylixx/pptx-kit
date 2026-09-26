# -*- coding: utf-8 -*-
"""export_thumbs.py 测试：
1) Pillow 兜底路径：2 页 PPTX → 导出 2 张 PNG（尺寸正确）
2) 文件名前缀规则正确
（COM/LibreOffice 路径依赖本机 Office，不纳入离线测试）
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from export_thumbs import _try_pillow

PASS = FAIL = 0


def ok(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        print(f"  ❌ {name} {detail}")


def make_mini_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for i in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
        tb.text = f"Page {i + 1}"
    prs.save(path)


tmp = tempfile.mkdtemp(prefix="ppt_thumb_test_")
pptx_path = os.path.join(tmp, "mini.pptx")
make_mini_pptx(pptx_path)

print("=== export_thumbs Pillow 路径 ===")
out_dir = os.path.join(tmp, "out")
r = _try_pillow(pptx_path, out_dir, 640, 360, "slide")
ok("Pillow 导出返回 True", r is True)
files = [f for f in os.listdir(out_dir) if f.endswith(".png")]
ok("生成 2 张 PNG", len(files) == 2, str(files))
ok("文件名前缀正确", sorted(files) == ["slide001.png", "slide002.png"], str(files))
from PIL import Image
im = Image.open(os.path.join(out_dir, "slide001.png"))
ok("尺寸 640x360", im.size == (640, 360), str(im.size))

print(f"\n结果：{PASS} 通过 / {FAIL} 失败")
sys.exit(1 if FAIL else 0)
