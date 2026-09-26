#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
scan_layouts.py — 模板新版式 → patterns.json 槽位映射的自动生成/更新工具。

用途：
  在 template.pptx 中新增了版式页后，不必手工数形状索引写 patterns.json。
  对新增页的槽位形状按约定命名（title/s1/b1/img1/图表/表格等），运行本工具，
  自动生成该版式的 patterns.json 条目，并可同步更新 layout_index.json 的
  intent 候选列表与 storyboard.schema.json 的 pattern enum。

用法（三种模式）：

1) 命名扫描（推荐）——先在 PowerPoint 里把新页的槽位形状重命名为槽位 key：
   python scan_layouts.py --page 71 --id text-3icon --desc "三卡片带图标" \
                          --intent points-3
   （默认写回 patterns.json；--intent 同时更新 layout_index.json 与 schema）

2) 快速映射——不重命名，用命令行直接给"形状索引=槽位key"：
   python scan_layouts.py --page 71 --id text-3icon --map "title=0,s1=1,b1=2,img1=3"

3) 交互映射——逐个形状询问：
   python scan_layouts.py --page 71 --id text-3icon --interactive

校验既有映射是否随模板改动漂移：
   python scan_layouts.py --check

全模板版式盘点（参考 pptx-profile：自动输出每页结构指纹、建议版式族与
intent，并对照 patterns.json 显示已接入/未接入）：
   python scan_layouts.py --inventory

其他选项：
   --template PATH   模板文件（默认 assets/template.pptx）
   --patterns PATH   patterns.json（默认同目录）
   --layout PATH     layout_index.json（默认同目录）
   --schema PATH     storyboard.schema.json（默认同目录）
   --force           覆盖已存在的 pattern id（默认冲突时报错）
   --dry-run         只打印结果，不写文件（写回前先备份 .bak）
"""
import argparse
import json
import os
import re
import shutil
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

# 模板引擎自动生成的形状默认名（中英混排）：一律忽略，不当作槽位。
# 覆盖"矩形 43"、"矩形: 圆角 3"、"组合 22"、"TextBox 1" 等形态：
#   前缀(形状名词) + 可选(:/空格 + 修饰词)* + 可选(空格 + 序号)
DEFAULT_NAME_RE = re.compile(
    r"^(矩形|圆角矩形|椭圆|组合|文本占位符|直接连接符|直线|连接符|图片|图表|"
    r"表格|任意多边形|菱形|心形|左箭头|右箭头|上箭头|下箭头|直角箭头|弯箭头|"
    r"环形箭头|五边形|六边形|八边形|云形|燕尾形箭头|双箭头|文本框|占位符|"
    r"TextBox|Picture|Table|Chart|Group|Oval|Rectangle|Rounded Rectangle|"
    r"Placeholder|Content Placeholder|Freeform|Line|Straight Connector|"
    r"Isosceles Triangle|Right Triangle|Diamond|Hexagon|Cloud|Arrow|"
    r"Block Arc|Frame|Shape)"
    r"([ :]+[\u4e00-\u9fffA-Za-z]+)*( \d+)?$")


def is_user_named(sh):
    """形状是否有用户自定义的语义命名（非引擎默认名）。"""
    return bool(sh.name) and not DEFAULT_NAME_RE.match(sh.name)


def _is_image(sh):
    """图片槽判断：图片占位符 / PICTURE 类型 / 含 a:blip 引用。"""
    if getattr(sh, "is_placeholder", False) \
            and sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
        return True
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    if sh._element.find(".//" + qn("a:blip")) is not None:
        return True
    return False


def shape_kinds(sh):
    """形状可承载的槽位类型列表。"""
    kinds = []
    if sh.has_text_frame:
        kinds.append("text")
    if sh.has_table:
        kinds.append("table")
    if sh.has_chart:
        kinds.append("chart")
    if _is_image(sh):
        kinds.append("image")
    return kinds


def shape_preview(sh, maxlen=30):
    if sh.has_text_frame and sh.text_frame.text.strip():
        return sh.text_frame.text.strip().replace("\n", "|")[:maxlen]
    return ""


def scan_named(slide):
    """命名扫描：只收录用户自定义命名的形状（按模板形状索引）。"""
    text_slots, image_slots, charts, table_idx = {}, {}, {}, None
    notes = []
    for idx, sh in enumerate(slide.shapes):
        if not is_user_named(sh):
            continue
        kinds = shape_kinds(sh)
        if not kinds:
            notes.append(f"  ⚠ 形状[{idx}] '{sh.name}' 无文本/图片/图表/表格能力，忽略")
            continue
        if "table" in kinds:
            table_idx = idx
        elif "chart" in kinds:
            charts[sh.name] = idx
        elif "image" in kinds:
            image_slots[sh.name] = idx
        elif "text" in kinds:
            text_slots[sh.name] = idx
        else:
            notes.append(f"  ⚠ 形状[{idx}] '{sh.name}' 类型未识别，忽略")
    return text_slots, image_slots, charts, table_idx, notes


def build_entry(slide_no, text_slots, image_slots, charts, table_idx, desc):
    entry = {"slide": slide_no, "desc": desc or ""}
    if text_slots:
        entry["text_slots"] = text_slots
    if image_slots:
        entry["image_slots"] = image_slots
    if charts:
        if len(charts) == 1:
            entry["chart_slot"] = next(iter(charts.values()))
        else:
            entry["chart_slots"] = charts
    if table_idx is not None:
        entry["table_slot"] = table_idx
    return entry


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_json(path, data, backup=True):
    if backup and os.path.exists(path):
        shutil.copy2(path, path + ".bak")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def merge_patterns(path, pid, entry, force):
    data = load_json(path)
    if pid in data and not force:
        raise SystemExit(f"patterns.json 已存在 {pid}；用 --force 覆盖，或换 --id")
    data[pid] = entry
    save_json(path, data)
    print(f"  ✅ patterns.json ← {pid}")


def merge_schema(path, pid):
    data = load_json(path)
    enum = data["properties"]["pages"]["items"]["properties"]["pattern"]["enum"]
    if pid not in enum:
        enum.append(pid)
        save_json(path, data)
        print(f"  ✅ storyboard.schema.json ← enum += {pid}")


def merge_layout(path, intent, pid, intent_desc):
    data = load_json(path)
    if intent in data:
        cands = data[intent].setdefault("candidates", [])
        if pid not in cands:
            cands.append(pid)
            save_json(path, data)
            print(f"  ✅ layout_index.json [{intent}].candidates += {pid}")
        else:
            print(f"  · layout_index.json [{intent}] 已含 {pid}，跳过")
    else:
        data[intent] = {"desc": intent_desc or "", "candidates": [pid]}
        save_json(path, data)
        print(f"  ✅ layout_index.json 新建 intent '{intent}' ← {pid}")


def run_check(patterns_path, template_path):
    patterns = {k: v for k, v in load_json(patterns_path).items()
                if not k.startswith("_")}
    prs = Presentation(template_path)
    n_slides = len(prs.slides)
    drift = 0
    print(f"=== 校验 {len(patterns)} 个版式映射 vs 模板 {n_slides} 页 ===")
    for pid, entry in patterns.items():
        slide_no = entry.get("slide")
        if not (1 <= slide_no <= n_slides):
            print(f"  ❌ {pid}: slide {slide_no} 越界（模板共 {n_slides} 页）")
            drift += 1
            continue
        shapes = list(prs.slides[slide_no - 1].shapes)
        for kind, slots in (("text", entry.get("text_slots", {})),
                            ("image", entry.get("image_slots", {}))):
            for key, idx in slots.items():
                if idx >= len(shapes):
                    print(f"  ❌ {pid}: {kind}_slots['{key}'] 索引 {idx} 越界（该页 {len(shapes)} 形状）")
                    drift += 1
                    continue
                sh = shapes[idx]
                expect = "text" if kind == "text" else "image"
                if expect == "text" and not sh.has_text_frame:
                    print(f"  ❌ {pid}: 索引 {idx}（{kind} '{key}'）不再是文本框，可能是 {sh.name}")
                    drift += 1
                elif expect == "image" and not _is_image(sh):
                    print(f"  ❌ {pid}: 索引 {idx}（{kind} '{key}'）不再是图片，可能是 {sh.name}")
                    drift += 1
        if entry.get("chart_slot") is not None:
            idx = entry["chart_slot"]
            if idx >= len(shapes) or not shapes[idx].has_chart:
                print(f"  ❌ {pid}: chart_slot {idx} 不再是图表")
                drift += 1
        for key, idx in entry.get("chart_slots", {}).items():
            if idx >= len(shapes) or not shapes[idx].has_chart:
                print(f"  ❌ {pid}: chart_slots['{key}'] {idx} 不再是图表")
                drift += 1
        if entry.get("table_slot") is not None:
            idx = entry["table_slot"]
            if idx >= len(shapes) or not shapes[idx].has_table:
                print(f"  ❌ {pid}: table_slot {idx} 不再是表格")
                drift += 1
    print(f"\n{'✅ 未发现漂移' if drift == 0 else f'❌ 发现 {drift} 处漂移'}")
    sys.exit(1 if drift else 0)


def page_fingerprint(slide):
    """提取单页版式的结构指纹：文本/图片/图表/表格数量与文本形态。"""
    n_text = n_img = n_chart = n_table = n_short = 0
    short_labels, first_lines = [], []
    for sh in slide.shapes:
        kinds = shape_kinds(sh)
        if "text" in kinds:
            n_text += 1
            t = shape_preview(sh, 16)
            if t:
                first_lines.append(t)
                if len(t) <= 10:
                    n_short += 1
                    short_labels.append(t)
        if "image" in kinds:
            n_img += 1
        if "chart" in kinds:
            n_chart += 1
        if "table" in kinds:
            n_table += 1
    return {
        "n_text": n_text, "n_img": n_img, "n_chart": n_chart,
        "n_table": n_table, "n_short": n_short,
        "short_labels": short_labels[:6], "first_lines": first_lines[:3],
    }


def suggest_family(fp):
    """启发式：由结构指纹推断版式族与建议 intent（借鉴 pptx-profile 的用途推断）。"""
    c = fp["n_chart"]
    if c >= 3:
        return "chart 多图", "chart", "多圆环/多图对比（chart_slots）"
    if c >= 1:
        return "chart-*", "chart", "单图表页（chart_slot）"
    if fp["n_table"] >= 1:
        return "table", "table", "数据表格页（table_slot）"
    if fp["n_img"] >= 2 and fp["n_short"] >= 2:
        return "imgtext-N / mockup", "points-3", "多图文卡片（image_slots）"
    if fp["n_img"] >= 1:
        return "imgtext / mockup", "points-3", "图文页（image_slot）"
    if fp["n_short"] >= 3:
        return f"text-{min(fp['n_short'], 6)} / steps-{min(fp['n_short'], 6)}", \
               f"points-{min(fp['n_short'], 6)}", f"{fp['n_short']} 组短标签要点"
    if fp["n_short"] == 2:
        return "text-2 / compare-2", "points-2", "双要点或对比页"
    if fp["n_text"] <= 1:
        return "cover / section / thanks", "cover", "单文本页（封面/章节/结尾）"
    return "通用内容页", "points-3", "结构不典型，需人工确认"


def run_inventory(template_path, patterns_path):
    """全模板版式盘点：结构指纹聚类 + 建议版式族 + 已接入对照。"""
    patterns = {k: v for k, v in load_json(patterns_path).items()
                if not k.startswith("_")}
    by_slide = {}
    for pid, entry in patterns.items():
        by_slide.setdefault(int(entry["slide"]), []).append(pid)

    prs = Presentation(template_path)
    fps = {i: page_fingerprint(slide)
           for i, slide in enumerate(prs.slides, 1)}

    # 版式族聚类：相同结构指纹的页归为一族
    fam_pages = {}
    for i, fp in fps.items():
        key = (fp["n_text"], fp["n_img"], fp["n_chart"], fp["n_table"],
               fp["n_short"])
        fam_pages.setdefault(key, []).append(i)
    families = sorted(fam_pages.items(),
                      key=lambda kv: (kv[0][0], kv[0][1]))

    print(f"=== 模板全盘点（{len(prs.slides)} 页，已接入 {len(patterns)} 版式） ===")
    print("\n-- 版式族聚类（相同结构指纹的页） --")
    for (n_t, n_i, n_c, n_tb, n_s), pages in families:
        wired = [p for p in pages if p in by_slide]
        free = [p for p in pages if p not in by_slide]
        if not free:
            continue  # 全部已接入的族不占用注意力
        fam, intent, note = suggest_family(fps[pages[0]])
        mark = f"已接入{len(wired)}页，未接入{len(free)}页：{','.join(map(str, free))}"
        print(f"  指纹(文本{n_t} 图{n_i} 图表{n_c} 表{n_tb}) → {fam:<20} "
              f"intent={intent:<9} 页[{','.join(map(str, pages))}] · {mark}")
        if note and note != "结构不典型，需人工确认":
            print(f"      {note}")

    print("\n-- 逐页明细 --")
    print("| 页 | 形状构成 | 建议版式族 | 建议intent | 已接入 patterns | 文本预览 |")
    print("|----|----------|------------|------------|------------------|----------|")
    for i, fp in sorted(fps.items()):
        fam, intent, _ = suggest_family(fp)
        wired = ",".join(by_slide.get(i, [])) or "—"
        preview = (fp["first_lines"][0] if fp["first_lines"] else "")
        print(f"| {i:<3} | 文本{fp['n_text']} 图{fp['n_img']} 图表{fp['n_chart']} "
              f"表{fp['n_table']} | {fam:<22} | {intent:<10} | {wired:<30} | "
              f"{preview} |")

    free_text = [i for i in range(1, len(prs.slides) + 1)
                 if i not in by_slide and fps[i]["n_text"] > 0]
    print(f"\n未接入且含文本的页（潜在可用版式）：{', '.join(map(str, free_text)) or '无'}")
    print("提示：使用说明页（如 55-70 的字体/图表教程）结构相似但不可当版式接入，"
          "请结合『文本预览』列判断；确认是版式后，用 --page N --map / "
          "--interactive 或命名扫描录入 patterns.json。")


def main():
    ap = argparse.ArgumentParser(
        description="模板新版式 → patterns.json 槽位映射自动生成/更新")
    ap.add_argument("--template", default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "assets", "template.pptx"))
    ap.add_argument("--patterns", default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "patterns.json"))
    ap.add_argument("--layout", default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "layout_index.json"))
    ap.add_argument("--schema", default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "storyboard.schema.json"))
    ap.add_argument("--page", type=int, help="新版式所在模板页码（1-based）")
    ap.add_argument("--id", help="新版式 id（如 text-3icon）")
    ap.add_argument("--desc", default="", help="版式描述（可选）")
    ap.add_argument("--intent", help="把新版式加入哪个 intent 的候选（如 points-3）")
    ap.add_argument("--intent-desc", default="", help="新建 intent 时的描述")
    ap.add_argument("--map", help="快速映射：逗号分隔 key=索引，如 title=0,s1=1,img1=3")
    ap.add_argument("--interactive", action="store_true", help="交互式逐个分配槽位")
    ap.add_argument("--force", action="store_true", help="覆盖已存在的 pattern id")
    ap.add_argument("--dry-run", action="store_true", help="只打印结果，不写文件")
    ap.add_argument("--check", action="store_true", help="校验既有映射是否漂移")
    ap.add_argument("--inventory", action="store_true",
                    help="全模板版式盘点：结构指纹 + 建议版式族 + 已接入对照")
    a = ap.parse_args()

    if a.inventory:
        run_inventory(a.template, a.patterns)
        return

    if a.check:
        run_check(a.patterns, a.template)
        return

    if not a.page:
        ap.error("需要 --page（或 --check）")
    if not a.id:
        ap.error("需要 --id（新版式名称）")
    prs = Presentation(a.template)
    if not 1 <= a.page <= len(prs.slides):
        ap.error(f"页码 {a.page} 越界（模板共 {len(prs.slides)} 页）")
    slide = prs.slides[a.page - 1]

    if a.map:
        # 快速映射：key=idx 列表
        text_slots, image_slots, charts, table_idx = {}, {}, {}, None
        for pair in a.map.split(","):
            if not pair.strip():
                continue
            key, _, idx_s = pair.strip().partition("=")
            idx = int(idx_s)
            sh = slide.shapes[idx]
            kinds = shape_kinds(sh)
            if "table" in kinds:
                table_idx = idx
            elif "chart" in kinds:
                charts[key] = idx
            elif "image" in kinds:
                image_slots[key] = idx
            elif "text" in kinds:
                text_slots[key] = idx
            else:
                print(f"  ⚠ 索引 {idx}（{sh.name}）无可承载槽位类型，跳过 {key}")
        entry = build_entry(a.page, text_slots, image_slots, charts,
                            table_idx, a.desc)
    elif a.interactive:
        text_slots, image_slots, charts, table_idx = {}, {}, {}, None
        print(f"=== 第 {a.page} 页形状清单（共 {len(slide.shapes)} 个） ===")
        for idx, sh in enumerate(slide.shapes):
            print(f"  [{idx}] {sh.name!r} {','.join(shape_kinds(sh)) or '无能力'} "
                  f"{shape_preview(sh)!r}")
        print("输入槽位 key（如 title/s1/img1），回车跳过，输入 q 结束：")
        for idx in range(len(slide.shapes)):
            key = input(f"  索引 {idx} → key: ").strip()
            if key.lower() == "q":
                break
            if not key:
                continue
            sh = slide.shapes[idx]
            kinds = shape_kinds(sh)
            if "table" in kinds:
                table_idx = idx
            elif "chart" in kinds:
                charts[key] = idx
            elif "image" in kinds:
                image_slots[key] = idx
            elif "text" in kinds:
                text_slots[key] = idx
            else:
                print(f"    ⚠ 该形状无可承载槽位类型，忽略")
        entry = build_entry(a.page, text_slots, image_slots, charts,
                            table_idx, a.desc)
    else:
        # 命名扫描
        text_slots, image_slots, charts, table_idx, notes = scan_named(slide)
        for n in notes:
            print(n)
        entry = build_entry(a.page, text_slots, image_slots, charts,
                            table_idx, a.desc)

    print(f"=== 生成条目（第 {a.page} 页，id={a.id}） ===")
    print(json.dumps(entry, ensure_ascii=False, indent=2))
    if not entry.get("text_slots") and not entry.get("image_slots") \
            and not entry.get("chart_slot") and not entry.get("chart_slots") \
            and entry.get("table_slot") is None:
        print("  ⚠ 未识别到任何槽位。若用了命名扫描，请先在 PowerPoint 的"
              "「选择窗格」里给槽位形状重命名（如 title/s1/img1）；"
              "或改用 --map / --interactive 手动指定。")
        sys.exit(1)

    if a.dry_run:
        print("\n(--dry-run：未写文件)")
        return

    merge_patterns(a.patterns, a.id, entry, a.force)
    merge_schema(a.schema, a.id)
    if a.intent:
        merge_layout(a.layout, a.intent, a.id, a.intent_desc)


if __name__ == "__main__":
    main()
