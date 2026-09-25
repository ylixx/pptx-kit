#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
render_pptx.py — 把 storyboard.json + 版式库模板渲染成原生可编辑 PPTX。

两阶段：
  1) 在 OOXML(zip) 层把模板中被选中的"单页版式"按 storyboard 顺序复制成新 deck；
  2) 用 python-pptx 按 patterns.json 的槽位映射填文字 / 换图 / 写图表表格 / 写备注。

用法：
  python3 render_pptx.py storyboard.json -o out.pptx
  python3 render_pptx.py storyboard.json -t <模板.pptx> -p <patterns.json> -o out.pptx
  python3 render_pptx.py --list-patterns

默认资产（模板/映射）按 skill 布局（../assets/）与独立布局（./）自动探测。

依赖：python-pptx (pip install python-pptx)
"""

import argparse
import json
import os
import posixpath
import re
import sys
import zipfile
from copy import deepcopy

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

from checks import page_issues

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_ROOT = os.path.dirname(HERE)


def _default_asset(name):
    """兼容 skill 布局(../assets/)与独立包布局(./ 或 ./assets/)的资产路径探测。"""
    for cand in (os.path.join(SKILL_ROOT, "assets", name),
                 os.path.join(HERE, name),
                 os.path.join(HERE, "assets", name)):
        if os.path.exists(cand):
            return cand
    return os.path.join(SKILL_ROOT, "assets", name)


SLIDE_CT = ("application/vnd.openxmlformats-officedocument"
            ".presentationml.slide+xml")
SLIDE_RT = ("http://schemas.openxmlformats.org/officeDocument"
            "/2006/relationships/slide")

WARNINGS = []


def warn(msg):
    WARNINGS.append(msg)
    print(f"[warn] {msg}", file=sys.stderr)


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        raw = f.read()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        # 网页大模型常把 JSON 包在 ```json 代码块或前后缀文字里，自动提取
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start:end + 1])
        raise


# ---------------------------------------------------------------- 阶段 1
def assemble_deck(template_path, donor_positions, out_path):
    """把模板中 donor_positions(1-based 版式页码, 可重复) 复制成顺序一致的新 deck。"""
    with zipfile.ZipFile(template_path) as z:
        entries = {n: z.read(n) for n in z.namelist()}

    pres = entries["ppt/presentation.xml"].decode("utf-8")
    rels = entries["ppt/_rels/presentation.xml.rels"].decode("utf-8")
    ctypes = entries["[Content_Types].xml"].decode("utf-8")

    # rId -> "slides/slideN.xml"
    rel_map = dict(re.findall(
        r'<Relationship [^>]*Id="(rId\d+)"[^>]*Target="(slides/slide\d+\.xml)"',
        rels))
    order = [rel_map[rid] for rid in re.findall(
        r'<p:sldId [^>]*r:id="(rId\d+)"', pres)]

    donors = []
    for pos in donor_positions:
        if not 1 <= pos <= len(order):
            raise SystemExit(f"模板页码越界: {pos} (模板共 {len(order)} 页)")
        donors.append(order[pos - 1])          # "slides/slideN.xml"

    k = len(donors)
    # 新 slide 部件（沿用 slide1..slideK 命名，rels 逐字复制并剥离 notesSlide）
    new_parts = {}
    for i, donor in enumerate(donors, 1):
        base = os.path.basename(donor)                      # slideN.xml
        src_rels = f"ppt/slides/_rels/{base}.rels"
        new_parts[f"ppt/slides/slide{i}.xml"] = entries[
            f"ppt/slides/{base}"]
        rels_bytes = entries.get(src_rels, b"")
        rels_bytes = re.sub(
            rb'<Relationship [^>]*Type="[^"]*/notesSlide"[^>]*/>', b"",
            rels_bytes)
        new_parts[f"ppt/slides/_rels/slide{i}.xml.rels"] = rels_bytes

    # 图表部件去深拷贝：同一图表被多页复用时，后到者拿到独立副本，
    # 否则 replace_data 会互相覆盖数据。
    part_ct = {m.group(1).lstrip("/"): m.group(2)
               for m in re.finditer(
                   r'<Override PartName="([^"]+)" ContentType="([^"]+)"',
                   ctypes)}
    chart_seen = {}          # 原始 chart 部件名 -> 已分配副本数
    extra_parts = {}
    ct_extra = []

    def clone_with_deps(part_abs, tag):
        """深拷贝一个部件（连同其 _rels 里的依赖），返回新部件绝对名。"""
        dirname, base = posixpath.split(part_abs)
        stem, ext = posixpath.splitext(base)
        new_abs = posixpath.join(dirname, f"{stem}_c{tag}{ext}")
        extra_parts[new_abs] = entries[part_abs]
        if part_abs in part_ct:            # 克隆件需注册 Override
            ct_extra.append(
                f'<Override PartName="/{new_abs}" '
                f'ContentType="{part_ct[part_abs]}"/>')
        rels_abs = posixpath.join(dirname, "_rels", base + ".rels")
        if rels_abs in entries:
            rels_txt = entries[rels_abs].decode("utf-8")

            def repl(m):
                tag_txt = m.group(0)
                if 'TargetMode="External"' in tag_txt:
                    return tag_txt
                tgt = re.search(r'Target="([^"]+)"', tag_txt).group(1)
                dep_abs = posixpath.normpath(posixpath.join(dirname, tgt))
                dep_dir, dep_base = posixpath.split(dep_abs)
                dep_stem, dep_ext = posixpath.splitext(dep_base)
                dep_new = posixpath.join(dep_dir, f"{dep_stem}_c{tag}{dep_ext}")
                if dep_abs in entries:
                    extra_parts[dep_new] = entries[dep_abs]
                    if dep_abs in part_ct:  # 依赖件同样注册 Override
                        ct_extra.append(
                            f'<Override PartName="/{dep_new}" '
                            f'ContentType="{part_ct[dep_abs]}"/>')
                new_tgt = posixpath.join(
                    posixpath.relpath(dep_dir, dirname),
                    f"{dep_stem}_c{tag}{dep_ext}")
                return tag_txt.replace(f'Target="{tgt}"', f'Target="{new_tgt}"')

            rels_txt = re.sub(r"<Relationship [^>]*/>", repl, rels_txt)
            extra_parts[posixpath.join(
                dirname, "_rels", f"{stem}_c{tag}{ext}.rels")] = \
                rels_txt.encode("utf-8")
        return new_abs

    for i in range(1, k + 1):
        rels_name = f"ppt/slides/_rels/slide{i}.xml.rels"
        rels_bytes = new_parts[rels_name]
        rels_txt = rels_bytes.decode("utf-8")

        def repl_slide(m):
            tag_txt = m.group(0)
            if 'relationships/chart"' not in tag_txt:
                return tag_txt
            tgt = re.search(r'Target="([^"]+)"', tag_txt).group(1)
            chart_abs = posixpath.normpath(posixpath.join("ppt/slides", tgt))
            seen = chart_seen.get(chart_abs, 0)
            chart_seen[chart_abs] = seen + 1
            if seen == 0:
                return tag_txt
            new_abs = clone_with_deps(chart_abs, seen)
            new_tgt = posixpath.join(
                posixpath.relpath(posixpath.dirname(new_abs), "ppt/slides"),
                posixpath.basename(new_abs))
            return tag_txt.replace(f'Target="{tgt}"', f'Target="{new_tgt}"')

        rels_txt = re.sub(r"<Relationship [^>]*/>", repl_slide, rels_txt)
        new_parts[rels_name] = rels_txt.encode("utf-8")

    new_parts.update(extra_parts)

    # [Content_Types].xml：清掉旧 slide Override，补新的
    ctypes = re.sub(
        r'<Override PartName="/ppt/slides/slide\d+\.xml"[^>]*/>', "", ctypes)
    adds = "".join(
        f'<Override PartName="/ppt/slides/slide{i}.xml" '
        f'ContentType="{SLIDE_CT}"/>' for i in range(1, k + 1))
    adds += "".join(ct_extra)
    ctypes = ctypes.replace("</Types>", adds + "</Types>")

    # presentation rels：清掉旧 slide 关系，补新关系（rId901+）
    rels = re.sub(
        r'<Relationship [^>]*Type="[^"]*/slide"[^>]*/>', "", rels)
    adds = "".join(
        f'<Relationship Id="rId{900 + i}" Type="{SLIDE_RT}" '
        f'Target="slides/slide{i}.xml"/>' for i in range(1, k + 1))
    rels = rels.replace("</Relationships>", adds + "</Relationships>")

    # presentation.xml：重写 sldIdLst
    lst = "".join(
        f'<p:sldId id="{256 + i}" r:id="rId{900 + i}"/>' for i in range(1, k + 1))
    pres, n = re.subn(r"<p:sldIdLst>.*?</p:sldIdLst>",
                      f"<p:sldIdLst>{lst}</p:sldIdLst>", pres, flags=re.S)
    if n != 1:
        raise SystemExit("presentation.xml 中未找到 sldIdLst")

    # 写出：丢弃全部旧 slide 部件，写入新部件
    drop = re.compile(r"^ppt/slides/slide\d+\.xml$")
    drop_rels = re.compile(r"^ppt/slides/_rels/slide\d+\.xml\.rels$")
    entries["ppt/presentation.xml"] = pres.encode("utf-8")
    entries["ppt/_rels/presentation.xml.rels"] = rels.encode("utf-8")
    entries["[Content_Types].xml"] = ctypes.encode("utf-8")

    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as z:
        final = {}
        for name, data in entries.items():
            if drop.match(name) or drop_rels.match(name):
                continue
            final[name] = data
        final.update(new_parts)
        # 裁剪未被引用的媒体素材（未选中的模板页图片），避免输出膨存
        referenced = set()
        for name, data in final.items():
            if not name.endswith(".rels"):
                continue
            base = posixpath.dirname(posixpath.dirname(name))
            for m in re.finditer(rb"<Relationship [^>]*/>", data):
                tag = m.group(0)
                if b'TargetMode="External"' in tag:
                    continue
                t = re.search(rb'Target="([^"]+)"', tag)
                if not t:
                    continue
                tgt = t.group(1).decode("utf-8")
                if tgt.startswith("http"):
                    continue
                referenced.add(posixpath.normpath(posixpath.join(base, tgt)))
        for name in list(final):
            if name.startswith("ppt/media/") and name not in referenced:
                del final[name]
        for name, data in final.items():
            z.writestr(name, data)
    return k


# ---------------------------------------------------------------- 阶段 2
def fill_textframe(tf, value):
    """把文字写进文本框，保留模板首个 run 的字符格式与段落格式。"""
    txBody = tf._txBody
    paras = txBody.findall(qn("a:p"))
    tmpl_pPr = tmpl_rPr = None
    if paras:
        p0 = paras[0]
        pPr = p0.find(qn("a:pPr"))
        if pPr is not None:
            tmpl_pPr = deepcopy(pPr)
        r0 = p0.find(qn("a:r"))
        if r0 is not None:
            rPr = r0.find(qn("a:rPr"))
            if rPr is not None:
                tmpl_rPr = deepcopy(rPr)
    for p in paras:
        txBody.remove(p)

    text = value if isinstance(value, str) else ("" if value is None else str(value))
    lines = text.split("\n") if text else [""]
    for line in lines:
        p = etree.SubElement(txBody, qn("a:p"))
        if tmpl_pPr is not None:
            p.append(deepcopy(tmpl_pPr))
        r = etree.SubElement(p, qn("a:r"))
        if tmpl_rPr is not None:
            r.append(deepcopy(tmpl_rPr))
        t = etree.SubElement(r, qn("a:t"))
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = line


def shape_at(slide, idx, pattern_id):
    shapes = list(slide.shapes)
    if not 0 <= idx < len(shapes):
        warn(f"[{pattern_id}] 形状索引 {idx} 越界（该页共 {len(shapes)} 个形状），跳过")
        return None
    return shapes[idx]


def fill_image(slide, shape, path, pattern_id):
    if not os.path.exists(path):
        warn(f"[{pattern_id}] 图片不存在: {path}，保留原占位符")
        return
    try:                                   # 空图片占位符
        shape.insert_picture(path)
        return
    except Exception:
        pass
    blip = shape._element.find(".//" + qn("a:blip"))
    if blip is None:
        warn(f"[{pattern_id}] 形状 {shape.name} 不是图片容器，跳过换图")
        return
    _, rId = slide.part.get_or_add_image_part(path)
    blip.set(qn("r:embed"), rId)


def fill_chart(shape, spec, pattern_id, prefix=""):
    try:
        cats = spec.get("categories") or []
        series = dict(spec.get("series") or {})
        # 长度一致性治理：错配时告警并截断到较短者，避免错位数据静默写入
        for name, vals in list(series.items()):
            if len(vals) != len(cats):
                n = min(len(vals), len(cats))
                warn(f"[{pattern_id}] 图表{prefix}系列 '{name}' 长度 {len(vals)} "
                     f"与 categories 长度 {len(cats)} 不一致，已截断到 {n}")
                series[name] = vals[:n]
        cd = CategoryChartData()
        cd.categories = cats
        for name, vals in series.items():
            cd.add_series(name, vals)
        shape.chart.replace_data(cd)
    except Exception as e:
        warn(f"[{pattern_id}] 图表数据写入失败: {e}")


def fill_table(shape, spec, pattern_id):
    try:
        table = shape.table
        rows = spec["rows"]
        if len(rows) > len(table.rows):
            warn(f"[{pattern_id}] 表格行数不足（模板 {len(table.rows)} 行），多余数据丢弃")
        if rows and max(len(r) for r in rows) > len(table.columns):
            warn(f"[{pattern_id}] 表格列数不足（模板 {len(table.columns)} 列），多余数据丢弃")
        # 先清空整表（避免模板示例数据残留），再写入提供的数据
        for r in range(len(table.rows)):
            for c in range(len(table.columns)):
                val = rows[r][c] if (r < len(rows) and c < len(rows[r])) else ""
                fill_textframe(table.cell(r, c).text_frame, val)
    except Exception as e:
        warn(f"[{pattern_id}] 表格数据写入失败: {e}")


def _fallback_pattern(patterns, page):
    """未知 pattern 时，按字段覆盖与图片槽匹配选一个兜底版式。

    优先"覆盖字段多"，其次"槽位贴近"；返回 pattern id 或 None（无任何可容纳页）。
    """
    fk = set((page.get("fields") or {}).keys())
    has_img = bool(page.get("images"))
    best, best_score = None, None
    for pid, pat in patterns.items():
        slots = set(pat.get("text_slots", {}))
        if has_img and not pat.get("image_slots"):
            continue
        covered = len(fk & slots)
        score = (covered, -abs(len(slots) - len(fk)))
        if best_score is None or score > best_score:
            best, best_score = pid, score
    return best


def render(storyboard, patterns, base_dir, template_path, out_path):
    pages = storyboard["pages"]

    # 校验 pattern 与槽位 key（未知 pattern 不中断：兜底替换或跳过）
    skipped = []
    for i, page in enumerate(pages, 1):
        pid = page.get("pattern")
        if pid not in patterns:
            fb = _fallback_pattern(patterns, page)
            if fb:
                warn(f"第 {i} 页 pattern 未知: {pid!r}，已兜底替换为字段兼容版式 '{fb}'")
                page["_intent"] = pid
                page["pattern"] = fb
                pid = fb
            else:
                warn(f"第 {i} 页 pattern 未知: {pid!r} 且无兜底版式，跳过该页")
                skipped.append(page)
                continue
        pat = patterns[pid]
        for msg in page_issues(page, pat):
            warn(f"第 {i} 页 [{pid}] {msg}")

    if skipped:
        pages = [p for p in pages if p not in skipped]
        storyboard = dict(storyboard)
        storyboard["pages"] = pages

    donors = [patterns[p["pattern"]]["slide"] for p in pages]
    assemble_deck(template_path, donors, out_path)

    prs = Presentation(out_path)
    if len(prs.slides) != len(pages):
        raise SystemExit("内部错误：幻灯片数量与 storyboard 不一致")

    for page, slide in zip(pages, prs.slides):
        pid = page["pattern"]
        pat = patterns[pid]
        fields = page.get("fields") or {}

        for key, idx in pat.get("text_slots", {}).items():
            shape = shape_at(slide, idx, pid)
            if shape is None:
                continue
            if not shape.has_text_frame:
                warn(f"[{pid}] 形状 {shape.name} 无文本框（槽位 {key}），跳过")
                continue
            fill_textframe(shape.text_frame, fields.get(key, ""))
            value = fields.get(key) or ""
            key_base = re.sub(r"\d+$", "", key)
            limit = TEXT_LIMITS.get(key_base) or TEXT_LIMITS.get(key_base.split("_")[-1])
            if limit and len(value) > limit:
                warn(f"[{pid}] 槽位 {key} 超长 {len(value)}>{limit} 字，可能溢出版式: {value[:20]}…")

        for key, idx in pat.get("image_slots", {}).items():
            src = (page.get("images") or {}).get(key)
            if not src:
                continue
            shape = shape_at(slide, idx, pid)
            if shape is None:
                continue
            path = src if os.path.isabs(src) else os.path.join(base_dir, src)
            fill_image(slide, shape, path, pid)

        if page.get("chart") is not None and "chart_slot" in pat:
            shape = shape_at(slide, pat["chart_slot"], pid)
            if shape is not None:
                fill_chart(shape, page["chart"], pid, "chart: ")

        charts = page.get("charts") or {}
        for key, idx in pat.get("chart_slots", {}).items():
            if charts.get(key) is None:
                continue
            shape = shape_at(slide, idx, pid)
            if shape is not None:
                fill_chart(shape, charts[key], pid, f"图表槽位 '{key}': ")

        if page.get("table") is not None and "table_slot" in pat:
            shape = shape_at(slide, pat["table_slot"], pid)
            if shape is not None:
                fill_table(shape, page["table"], pid)

        if page.get("notes"):
            slide.notes_slide.notes_text_frame.text = page["notes"]

    prs.save(out_path)
    return len(pages)


# 字数上限（只用于告警，见 PROMPT.md）
TEXT_LIMITS = {
    "title": 16, "t": 16,
    "s": 8, "st": 8, "kw": 8, "lab": 8, "mid": 8,
    "b": 36, "desc": 36, "top": 36, "bot": 36,
    "body": 160, "bio": 80, "caption": 80, "center": 80,
    "en": 40, "d": 6, "pct": 6, "ref": 80, "cap": 12,
    "num": 4, "name": 12, "main": 40, "org": 30, "sub": 40,
    "presenter": 24, "advisor": 24, "source": 16, "table_title": 30,
    "major": 20, "edu": 20, "skill": 20, "work": 20, "honor": 20,
    "contact": 20, "award": 10, "event": 10,
}


def main():
    ap = argparse.ArgumentParser(description="storyboard.json + 版式库模板 → PPTX")
    ap.add_argument("storyboard", nargs="?", help="storyboard.json 路径")
    ap.add_argument("-t", "--template",
                    default=_default_asset("template.pptx"),
                    help="版式库模板 pptx（默认 assets/template.pptx）")
    ap.add_argument("-p", "--patterns",
                    default=_default_asset("patterns.json"),
                    help="槽位映射文件（默认 patterns.json）")
    ap.add_argument("-o", "--out", help="输出 pptx 路径")
    ap.add_argument("--list-patterns", action="store_true",
                    help="列出全部版式与槽位后退出")
    args = ap.parse_args()

    patterns = load_json(args.patterns)
    patterns = {k: v for k, v in patterns.items() if not k.startswith("_")}

    if args.list_patterns:
        for pid, pat in patterns.items():
            keys = ", ".join(pat.get("text_slots", {}))
            imgs = ", ".join(pat.get("image_slots", {}))
            print(f"{pid:<15} 模板第{pat['slide']:>2}页  {pat.get('desc','')}")
            print(f"{'':<15} text: {keys}")
            if imgs:
                print(f"{'':<15} image: {imgs}")
        return

    if not args.storyboard or not args.out:
        ap.error("需要 storyboard.json 与 -o 输出路径")

    storyboard = load_json(args.storyboard)

    # 意图层集成（向后兼容）：layout_index.json 存在时，把 storyboard 中的
    # intent 解析为具体 pattern；不存在则行为与原来完全一致。只做最小侵入。
    _li_path = _default_asset("layout_index.json")
    if os.path.exists(_li_path):
        sys.path.insert(0, HERE)
        from resolve_intents import resolve_intents
        _li = {k: v for k, v in load_json(_li_path).items()
               if not k.startswith("_")}
        storyboard = resolve_intents(storyboard, _li, patterns)

    base_dir = os.path.dirname(os.path.abspath(args.storyboard))
    n = render(storyboard, patterns, base_dir, args.template, args.out)
    print(f"✅ 已生成 {args.out}（{n} 页，{len(WARNINGS)} 条警告）")


if __name__ == "__main__":
    main()
